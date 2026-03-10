from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
import os

doc = Document()

# --- Page setup ---
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)
font.color.rgb = RGBColor(0x33, 0x33, 0x33)

# ============================================================
# TITLE PAGE
# ============================================================
for _ in range(6):
    doc.add_paragraph('')

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('Part 02\n')
run.bold = True
run.font.size = Pt(28)
run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

run = title.add_run('Technical Report\n')
run.bold = True
run.font.size = Pt(24)
run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

run = title.add_run('Server Bug Fixes & Deployment Issues\n')
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

run = title.add_run('\nتقرير الأخطاء التقنية — الجزء الثاني\n')
run.bold = True
run.font.size = Pt(20)
run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

run = title.add_run('إصلاحات أخطاء السيرفر ومشاكل النشر')
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_paragraph('')
doc.add_paragraph('')

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
for label, value in [
    ('Website / الموقع: ', 'https://mixjo.top'),
    ('\nProject / المشروع: ', 'tourism-management-system'),
    ('\nDate / التاريخ: ', '14 February 2026 — 14 فبراير 2026'),
    ('\nPrepared for / مُعد إلى: ', 'Mr. Islam (Lead Developer)'),
]:
    run = info.add_run(label)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)
    run = info.add_run(value)
    run.font.size = Pt(11)

doc.add_page_break()

# ============================================================
# TABLE OF CONTENTS
# ============================================================
h = doc.add_heading('Table of Contents / جدول المحتويات', level=1)
for run in h.runs:
    run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

toc_items = [
    '1. Executive Summary / ملخص تنفيذي',
    '2. Bug #1: Case Sensitivity — EntryPoints Controllers',
    '   خطأ #1: حساسية حالة الأحرف — وحدة نقاط الدخول',
    '3. Bug #2: Wrong Route Names — Roles Module',
    '   خطأ #2: أسماء Routes خاطئة — وحدة الأدوار',
    '4. Bug #3: Missing Database Tables on Server',
    '   خطأ #3: جداول قاعدة بيانات مفقودة',
    '5. Bug #4: Missing Buttons — Transportation Companies',
    '   خطأ #4: أزرار مفقودة — شركات النقل',
    '6. Bug #5: Export Format Issue',
    '   خطأ #5: مشكلة صيغة التصدير',
    '7. Summary of Modified Files / ملخص الملفات المعدلة',
    '8. Required Server Actions / الإجراءات المطلوبة',
    '9. Recommendations / التوصيات',
]
for item in toc_items:
    p = doc.add_paragraph(item)
    p.paragraph_format.space_after = Pt(2)
    if item.startswith('   '):
        p.paragraph_format.left_indent = Cm(1)
        for run in p.runs:
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

doc.add_page_break()

# ============================================================
# HELPER FUNCTIONS
# ============================================================
def add_section_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)
    return h

def add_label_value(label, value, bold_value=False):
    p = doc.add_paragraph()
    run = p.add_run(label)
    run.bold = True
    run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)
    run = p.add_run(value)
    if bold_value:
        run.bold = True
    return p

def add_code_block(code_text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x80)
    return p

def add_diff_block(old_lines, new_lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    for line in old_lines:
        run = p.add_run('- ' + line + '\n')
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
    for line in new_lines:
        run = p.add_run('+ ' + line + '\n')
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x00, 0x80, 0x00)
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header row
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        run = p.add_run(header)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        from docx.oxml.ns import qn
        shading = cell._element.get_or_add_tcPr()
        shading_elem = shading.makeelement(qn('w:shd'), {
            qn('w:val'): 'clear',
            qn('w:color'): 'auto',
            qn('w:fill'): '8B0000'
        })
        shading.append(shading_elem)
    # Data rows
    for r, row_data in enumerate(rows):
        for c, cell_text in enumerate(row_data):
            cell = table.rows[r + 1].cells[c]
            cell.text = str(cell_text)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
    return table

def add_warning_box(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run('⚠ ' + text)
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0xCC, 0x66, 0x00)

# ============================================================
# 1. EXECUTIVE SUMMARY
# ============================================================
add_section_heading('1. Executive Summary / ملخص تنفيذي')

doc.add_paragraph(
    'Five (5) major technical bugs were discovered on the live server (mixjo.top) '
    'that affect the functionality of the dashboard. Most of these bugs work normally '
    'on the local development environment (Windows) but fail on the production server '
    '(Linux) due to operating system differences and deployment gaps.'
)
doc.add_paragraph(
    'تم اكتشاف 5 أخطاء تقنية رئيسية على السيرفر الفعلي (mixjo.top) تؤثر على عمل لوحة التحكم. '
    'معظم هذه الأخطاء تعمل بشكل طبيعي على بيئة التطوير المحلية (Windows) لكنها تتعطل على '
    'السيرفر (Linux) بسبب اختلافات أنظمة التشغيل وفجوات في عملية النشر.'
)

doc.add_paragraph('')
add_table(
    ['#', 'Bug / الخطأ', 'Severity / الخطورة', 'Status / الحالة'],
    [
        ['1', 'Case Sensitivity — EntryPoints\nحساسية حالة الأحرف', '🔴 Critical / حرج', '✅ Fixed / تم الإصلاح'],
        ['2', 'Wrong Route Names — Roles\nأسماء Routes خاطئة', '🔴 Critical / حرج', '✅ Fixed / تم الإصلاح'],
        ['3', 'Missing Database Tables\nجداول مفقودة', '🟡 Medium / متوسط', '⚠ Needs server action'],
        ['4', 'Missing Buttons — Transport\nأزرار مفقودة', '🟡 Medium / متوسط', '⚠ Needs permissions fix'],
        ['5', 'Export Format Issue\nمشكلة صيغة التصدير', '🟢 Low / منخفض', '⚠ Needs server check'],
    ]
)

doc.add_page_break()

# ============================================================
# 2. BUG #1
# ============================================================
add_section_heading('2. Bug #1: Case Sensitivity — EntryPoints Controllers')
add_section_heading('خطأ #1: حساسية حالة الأحرف — وحدة نقاط الدخول', level=2)

add_label_value('Severity / الخطورة: ', '🔴 CRITICAL — Breaks entire dashboard / يُعطّل كامل لوحة التحكم', True)

add_label_value('Error Message / رسالة الخطأ:', '')
add_code_block('ReflectionException: Class "Modules\\EntryPoints\\Http\\Controllers\\LandCrossingController" does not exist')

add_section_heading('Root Cause / السبب:', level=3)
doc.add_paragraph(
    'The routes file (web.php) imports controller classes with different letter casing '
    'than the actual file and class names.'
)
doc.add_paragraph(
    'ملف Routes يستورد الكلاسات بأسماء مختلفة عن الأسماء الحقيقية للملفات والكلاسات.'
)

add_table(
    ['Import in Routes (Wrong ❌)', 'Actual File/Class Name (Correct ✅)'],
    [
        ['LandCrossingController (capital C)', 'LandcrossingController (lowercase c)'],
        ['SeaPortController (capital P)', 'SeaportController (lowercase p)'],
    ]
)

doc.add_paragraph('')
add_section_heading('Why it works locally but not on server?', level=3)
add_section_heading('لماذا يعمل محلياً ولا يعمل على السيرفر؟', level=3)

doc.add_paragraph('• Windows: File system is case-INSENSITIVE → LandCrossing = Landcrossing ✅')
doc.add_paragraph('  ويندوز: نظام الملفات غير حساس لحالة الأحرف')
doc.add_paragraph('• Linux (Server): File system is case-SENSITIVE → LandCrossing ≠ Landcrossing ❌')
doc.add_paragraph('  لينكس (السيرفر): نظام الملفات حساس لحالة الأحرف')

add_section_heading('Impact / التأثير:', level=3)
doc.add_paragraph(
    'This bug breaks ALL dashboard links, not just EntryPoints pages. '
    'When Laravel fails to load routes at startup, the entire application crashes with 500 error.'
)
doc.add_paragraph(
    'هذا الخطأ يكسر كل الروابط في لوحة التحكم وليس فقط صفحات نقاط الدخول. '
    'عندما يفشل Laravel في تحميل Routes عند بدء التشغيل، يتعطل التطبيق بالكامل بخطأ 500.'
)

add_section_heading('Fixed File / الملف المعدل:', level=3)
add_code_block('Modules/EntryPoints/Routes/web.php')

add_section_heading('Changes / التعديلات:', level=3)
add_diff_block(
    [
        'use Modules\\EntryPoints\\Http\\Controllers\\LandCrossingController;',
        'use Modules\\EntryPoints\\Http\\Controllers\\SeaPortController;',
        "Route::resource('land-crossings', LandCrossingController::class);",
        "Route::resource('seaports', SeaPortController::class);",
    ],
    [
        'use Modules\\EntryPoints\\Http\\Controllers\\LandcrossingController;',
        'use Modules\\EntryPoints\\Http\\Controllers\\SeaportController;',
        "Route::resource('land-crossings', LandcrossingController::class);",
        "Route::resource('seaports', SeaportController::class);",
    ]
)

add_warning_box(
    'Recommendation: Always ensure file names exactly match class names (PSR-4 Standard). '
    'Consider renaming both file AND class to use consistent PascalCase like LandCrossingController.\n'
    'توصية: يجب دائماً التأكد من تطابق أسماء الملفات مع أسماء الكلاسات بالضبط (معيار PSR-4).'
)

doc.add_page_break()

# ============================================================
# 3. BUG #2
# ============================================================
add_section_heading('3. Bug #2: Wrong Route Names — Roles Module')
add_section_heading('خطأ #2: أسماء Routes خاطئة — وحدة الأدوار', level=2)

add_label_value('Severity / الخطورة: ', '🔴 CRITICAL — Roles edit page returns 500 error', True)

add_label_value('Error Message / رسالة الخطأ:', '')
add_code_block(
    'Route [roles.index] not defined.\n'
    '(View: Modules/Core/Resources/views/roles/edit.blade.php)'
)

add_section_heading('Root Cause / السبب:', level=3)
doc.add_paragraph(
    'The code uses short route names (roles.index, roles.edit, roles.show) '
    'instead of the full prefixed names (dashboard.core.roles.index, etc.) '
    'registered in the route definitions.'
)
doc.add_paragraph(
    'الكود يستخدم أسماء routes مختصرة بدلاً من الأسماء الكاملة المسجلة في تعريفات Routes.'
)

add_section_heading('Fixed Files / الملفات المعدلة:', level=3)

# File 1
add_section_heading('File 1: Modules/Core/Http/Controllers/RoleController.php', level=4)
doc.add_paragraph('3 changes — store(), update(), and destroy() methods:')
add_diff_block(
    ["redirect()->route('roles.index')  // in store(), update(), destroy()"],
    ["redirect()->route('dashboard.core.roles.index')  // corrected"]
)

# File 2
add_section_heading('File 2: Modules/Core/Resources/views/roles/show.blade.php', level=4)
doc.add_paragraph('3 changes — lines 18, 23, 84:')
add_diff_block(
    [
        "route('roles.edit', $role->id)  // lines 18, 84",
        "route('roles.index')             // line 23",
    ],
    [
        "route('dashboard.core.roles.edit', $role->id)",
        "route('dashboard.core.roles.index')",
    ]
)

# File 3
add_section_heading('File 3: Modules/Core/Resources/views/livewire/roles1.blade.php', level=4)
doc.add_paragraph('2 changes — lines 33, 38:')
add_diff_block(
    [
        "route('roles.show', $role->id)  // line 33",
        "route('roles.edit', $role->id)  // line 38",
    ],
    [
        "route('dashboard.core.roles.show', $role->id)",
        "route('dashboard.core.roles.edit', $role->id)",
    ]
)

add_warning_box(
    'Recommendation: Review ALL modules to ensure full prefixed route names are used consistently.\n'
    'توصية: مراجعة جميع الوحدات للتأكد من استخدام أسماء Routes الكاملة بشكل متسق.'
)

doc.add_page_break()

# ============================================================
# 4. BUG #3
# ============================================================
add_section_heading('4. Bug #3: Missing Database Tables on Server')
add_section_heading('خطأ #3: جداول قاعدة بيانات مفقودة على السيرفر', level=2)

add_label_value('Severity / الخطورة: ', '🟡 Medium / متوسط', True)

add_label_value('Error Messages / رسائل الخطأ:', '')
add_code_block(
    '1. SQLSTATE[HY000]: no such table: entry_points\n'
    '2. SQLSTATE[42S02]: Table \'mixtop_mixjo_top.transportations_companies\' doesn\'t exist'
)

add_section_heading('Root Cause / السبب:', level=3)
doc.add_paragraph(
    'Database migrations for EntryPoints and Transportation modules were not executed on the production server.'
)
doc.add_paragraph(
    'لم يتم تشغيل Migrations لوحدتي EntryPoints و Transportation على سيرفر الإنتاج.'
)

add_section_heading('Solution / الحل:', level=3)
add_code_block(
    'cd /home/mixtop/project_root\n'
    'php artisan migrate --path=Modules/EntryPoints/Database/Migrations\n'
    'php artisan migrate --path=Modules/Transportation/Database/Migrations'
)

add_warning_box(
    'Recommendation: Create a deployment script that automatically runs php artisan migrate on each deploy.\n'
    'توصية: إنشاء سكريبت نشر يشغل تلقائياً php artisan migrate عند كل عملية نشر.'
)

doc.add_page_break()

# ============================================================
# 5. BUG #4
# ============================================================
add_section_heading('5. Bug #4: Missing Buttons — Transportation Companies')
add_section_heading('خطأ #4: أزرار مفقودة — صفحة شركات النقل', level=2)

add_label_value('Severity / الخطورة: ', '🟡 Medium / متوسط', True)
add_label_value('Page / الصفحة: ', '/dashboard/transportation/companies')

add_section_heading('Description / الوصف:', level=3)
doc.add_paragraph(
    'Show, Edit, and Delete buttons are missing from the companies table. '
    'Only Force Delete button is visible.'
)
doc.add_paragraph(
    'أزرار العرض والتعديل والحذف غير ظاهرة في جدول الشركات. يظهر فقط زر الحذف النهائي (Force Delete).'
)

add_section_heading('Root Cause / السبب:', level=3)
doc.add_paragraph(
    'The data-table.blade.php component uses permission checks to show/hide buttons. '
    'The CompanyPolicy.php defines these rules:'
)

add_table(
    ['Button / الزر', 'Permission Check', 'Status / الحالة'],
    [
        ['👁 Show', "user->can('manage_companies')", '❌ Missing'],
        ['✏️ Edit', "user->can('manage_companies')", '❌ Missing'],
        ['🗑 Delete', "user->can('manage_companies')", '❌ Missing'],
        ['💀 Force Delete', "user->hasRole('superadmin')", '✅ Works'],
    ]
)

doc.add_paragraph('')
doc.add_paragraph(
    'The user has the superadmin role but does NOT have the manage_companies permission assigned. '
    'Force Delete works because it checks the Role directly, not a Permission.'
)
doc.add_paragraph(
    'المستخدم لديه دور superadmin لكن ليس لديه صلاحية manage_companies. '
    'زر الحذف النهائي يعمل لأنه يتحقق من الدور مباشرة وليس من الصلاحية.'
)

add_section_heading('Solution / الحل:', level=3)
add_code_block(
    'php artisan tinker\n'
    '>>> $role = \\Spatie\\Permission\\Models\\Role::findByName(\'superadmin\');\n'
    '>>> $role->givePermissionTo(\'manage_companies\');'
)
doc.add_paragraph('Or from the dashboard: /dashboard/core/roles → superadmin → add manage_companies permission')
doc.add_paragraph('أو من لوحة التحكم: /dashboard/core/roles → superadmin → أضف صلاحية manage_companies')

doc.add_page_break()

# ============================================================
# 6. BUG #5
# ============================================================
add_section_heading('6. Bug #5: Export Format Issue')
add_section_heading('خطأ #5: مشكلة صيغة التصدير', level=2)

add_label_value('Severity / الخطورة: ', '🟢 Low / منخفض', True)

add_section_heading('Description / الوصف:', level=3)
doc.add_paragraph(
    'When exporting data as XLSX, the resulting file may not be a proper Excel file. '
    'It could be a CSV file with an .xlsx extension.'
)
doc.add_paragraph(
    'عند تصدير البيانات بصيغة XLSX، الملف الناتج قد لا يكون بصيغة Excel حقيقية. '
    'قد يكون ملف CSV بامتداد .xlsx.'
)

add_section_heading('Root Cause / السبب:', level=3)
doc.add_paragraph(
    'The system uses Spatie\\SimpleExcel\\SimpleExcelWriter which depends on the openspout library. '
    'If the library is not properly installed on the server, it may fall back to CSV format.'
)

add_section_heading('Solution / الحل:', level=3)
add_code_block(
    'cd /home/mixtop/project_root\n'
    'composer require openspout/openspout'
)

doc.add_page_break()

# ============================================================
# 7. SUMMARY OF MODIFIED FILES
# ============================================================
add_section_heading('7. Summary of Modified Files / ملخص الملفات المعدلة')

add_table(
    ['#', 'File Path / مسار الملف', 'Changes / التعديلات'],
    [
        ['1', 'Modules/EntryPoints/Routes/web.php', '4 changes — case sensitivity fix'],
        ['2', 'Modules/Core/Http/Controllers/RoleController.php', '3 changes — route names fix'],
        ['3', 'Modules/Core/Resources/views/roles/show.blade.php', '3 changes — route names fix'],
        ['4', 'Modules/Core/Resources/views/livewire/roles1.blade.php', '2 changes — route names fix'],
    ]
)

doc.add_paragraph('')
p = doc.add_paragraph()
run = p.add_run('Total: 12 changes across 4 files')
run.bold = True
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

p = doc.add_paragraph()
run = p.add_run('المجموع: 12 تعديل في 4 ملفات')
run.bold = True
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

doc.add_page_break()

# ============================================================
# 8. REQUIRED SERVER ACTIONS
# ============================================================
add_section_heading('8. Required Server Actions / الإجراءات المطلوبة على السيرفر')

doc.add_paragraph('Execute the following commands on the production server:')
doc.add_paragraph('نفّذ الأوامر التالية على سيرفر الإنتاج:')

add_code_block(
    '# Step 1: Upload the 4 modified files listed above\n'
    '# الخطوة 1: ارفع الملفات الأربعة المعدلة أعلاه\n\n'
    '# Step 2: Run database migrations\n'
    '# الخطوة 2: شغّل Migrations\n'
    'cd /home/mixtop/project_root\n'
    'php artisan migrate --path=Modules/EntryPoints/Database/Migrations\n'
    'php artisan migrate --path=Modules/Transportation/Database/Migrations\n\n'
    '# Step 3: Clear all caches\n'
    '# الخطوة 3: امسح جميع الكاش\n'
    'php artisan optimize:clear\n\n'
    '# Step 4: Fix permissions (add manage_companies to superadmin)\n'
    '# الخطوة 4: أصلح الصلاحيات\n'
    'php artisan tinker\n'
    '>>> $role = \\Spatie\\Permission\\Models\\Role::findByName(\'superadmin\');\n'
    '>>> $role->givePermissionTo(\'manage_companies\');\n'
    '>>> exit\n\n'
    '# Step 5: Install Excel library (optional)\n'
    '# الخطوة 5: ثبّت مكتبة Excel (اختياري)\n'
    'composer require openspout/openspout'
)

doc.add_page_break()

# ============================================================
# 9. RECOMMENDATIONS
# ============================================================
add_section_heading('9. General Recommendations / التوصيات العامة')

recommendations = [
    (
        '1. Case Sensitivity Audit / فحص حساسية الأحرف',
        'Review ALL route files across ALL modules to ensure class names match file names exactly (PSR-4 Standard). '
        'A single mismatch can crash the entire application on Linux.\n'
        'مراجعة جميع ملفات Routes في كل الوحدات للتأكد من تطابق أسماء الكلاسات مع أسماء الملفات بالضبط. '
        'خطأ واحد يمكن أن يُعطّل التطبيق بالكامل على Linux.'
    ),
    (
        '2. Route Names Audit / فحص أسماء Routes',
        'Review all code to ensure full prefixed route names (dashboard.module.resource.action) are used '
        'instead of short names. The EntryPoints controllers also use wrong route names (entrypoints.index).\n'
        'مراجعة جميع الكود للتأكد من استخدام الأسماء الكاملة للـ Routes بدلاً من الأسماء المختصرة.'
    ),
    (
        '3. Deployment Script / سكريبت النشر',
        'Create an automated deployment script that runs: php artisan migrate && php artisan optimize:clear '
        'This prevents missing tables and stale cache issues.\n'
        'إنشاء سكريبت نشر آلي يشغل migrate و optimize:clear لمنع مشاكل الجداول المفقودة والكاش القديم.'
    ),
    (
        '4. Linux Testing Environment / بيئة اختبار Linux',
        'Test the application on a Linux environment (e.g., Docker, WSL) before deploying to production '
        'to catch case-sensitivity and other OS-specific issues.\n'
        'اختبار التطبيق على بيئة Linux قبل النشر لاكتشاف مشاكل حساسية الأحرف ومشاكل نظام التشغيل.'
    ),
    (
        '5. Permission Seeder / بذر الصلاحيات',
        'Create a seeder that automatically assigns all necessary permissions to the superadmin role. '
        'This ensures new permissions are always granted after deployment.\n'
        'إنشاء seeder يعيّن تلقائياً جميع الصلاحيات اللازمة لدور superadmin عند كل عملية نشر.'
    ),
]

for title, desc in recommendations:
    add_section_heading(title, level=3)
    doc.add_paragraph(desc)

# ============================================================
# SAVE
# ============================================================
output_dir = r'G:\Report Mixjo Top\technical_report'
os.makedirs(output_dir, exist_ok=True)
output_path = os.path.join(output_dir, 'Part 02 Technical_Report_User_Feedback.docx')
doc.save(output_path)

print(f'\n=== Report created successfully! ===')
print(f'File: {output_path}')
print('Done!')
